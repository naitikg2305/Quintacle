from quintacle_core.nodes.kb.spreadsheet import read_spreadsheet, write_spreadsheet


def test_write_then_read_round_trips_values():
    data = write_spreadsheet([{"name": "Budget", "rows": [["Item", "Cost"], ["Rent", "1500"], ["Food", "400.5"]]}])

    result = read_spreadsheet(data)

    assert result["sheets"][0]["name"] == "Budget"
    assert result["sheets"][0]["rows"] == [["Item", "Cost"], ["Rent", "1500"], ["Food", "400.5"]]


def test_numeric_strings_become_real_numbers_on_write():
    data = write_spreadsheet([{"name": "Sheet1", "rows": [["100", "abc", "3.5"]]}])

    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(data))
    row = next(wb.active.iter_rows(values_only=True))
    assert row == (100, "abc", 3.5)


def test_multiple_sheets_round_trip():
    data = write_spreadsheet(
        [
            {"name": "One", "rows": [["a"]]},
            {"name": "Two", "rows": [["b"]]},
        ]
    )

    result = read_spreadsheet(data)

    assert [s["name"] for s in result["sheets"]] == ["One", "Two"]
