from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook


def read_spreadsheet(data: bytes) -> dict:
    """Every sheet as a plain grid of strings — good enough for "see my data
    and fix a cell," not a full-fidelity spreadsheet (formulas are read as
    their last-computed value, not re-evaluated; formatting is dropped)."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [["" if cell is None else str(cell) for cell in row] for row in ws.iter_rows(values_only=True)]
        sheets.append({"name": name, "rows": rows})
    return {"sheets": sheets}


def write_spreadsheet(sheets: list[dict]) -> bytes:
    """The inverse of read_spreadsheet() — rebuilds a fresh .xlsx from
    plain string grids. Values that look numeric are written as real numbers
    (so SUM/formatting in whatever opens it next still works), everything
    else stays text."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(title=sheet["name"][:31] or "Sheet")
        for row in sheet["rows"]:
            ws.append([_coerce(cell) for cell in row])
    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _coerce(value: str):
    if value == "":
        return None
    try:
        return int(value)
    except ValueError:
        pass
    try:
        return float(value)
    except ValueError:
        return value
